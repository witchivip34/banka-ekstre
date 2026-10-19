"""Çıktı profilleri: csv, xlsx, f1f6 ve CLI."""
from __future__ import annotations

import csv
import json

import pytest
import xlrd
from openpyxl import load_workbook

from banka_ekstre import Islem, ekstre_oku
from banka_ekstre.cikti import SUTUNLAR
from banka_ekstre.cikti import bordro_f1f6
from banka_ekstre.cikti import csv as csv_cikti
from banka_ekstre.cikti import xlsx as xlsx_cikti
from banka_ekstre.cli import main
from conftest import BANKALAR, uret


def _islem(tutar: float, n: int = 0) -> Islem:
    return Islem(f"{n + 1:02d}.01.2026", f"açıklama {n}", f"R{n}", tutar, None, "ziraat", "x.xlsx", n + 2)


# ─────────────────────────────── csv / xlsx ───────────────────────────────

def test_csv_yaz_geri_oku(ornek_yolu, tmp_path):
    e = ekstre_oku(ornek_yolu("ziraat"))
    yol = csv_cikti.yaz(e.islemler, tmp_path / "z.csv")
    ham = yol.read_bytes()
    assert ham.startswith(b"\xef\xbb\xbf")  # UTF-8 BOM
    with yol.open(encoding="utf-8-sig", newline="") as f:
        satirlar = list(csv.reader(f, delimiter=";"))
    assert satirlar[0] == SUTUNLAR
    assert len(satirlar) == len(e.islemler) + 1
    for satir, islem in zip(satirlar[1:], e.islemler):
        assert satir[0] == islem.tarih
        giren = float(satir[4].replace(",", ".")) if satir[4] else 0.0
        cikan = float(satir[5].replace(",", ".")) if satir[5] else 0.0
        assert giren - cikan == pytest.approx(islem.tutar)
        assert satir[7] == islem.kaynak_dosya


def test_xlsx_yaz_geri_oku(ornek_yolu, tmp_path):
    e = ekstre_oku(ornek_yolu("vakifbank"))
    yol = xlsx_cikti.yaz(e.islemler, tmp_path / "v.xlsx")
    ws = load_workbook(yol).active
    assert [c.value for c in ws[1]] == SUTUNLAR
    assert all(c.font.bold for c in ws[1])
    assert ws.max_row == len(e.islemler) + 1
    for r, islem in enumerate(e.islemler, start=2):
        assert ws.cell(r, 1).value == islem.tarih
        assert (ws.cell(r, 5).value or 0) - (ws.cell(r, 6).value or 0) == pytest.approx(islem.tutar)
        assert ws.cell(r, 5).number_format == "#,##0.00"
        assert ws.cell(r, 7).value == pytest.approx(islem.bakiye)


# ─────────────────────────────── f1f6 ───────────────────────────────

# Varsayılan profilde beklenen yerleşim: (giren kolonu, çıkan kolonu, çıkan negatif mi)
BEKLENEN_YERLESIM = {
    "ziraat": (4, 5, False),
    "yapikredi": (4, 5, True),
    "turkiyefinans": (4, 5, True),
    "garanti": (3, 4, False),
    "isbankasi": (2, 3, False),
    "vakifbank": (2, 3, True),
}


@pytest.mark.parametrize("banka", BANKALAR)
def test_f1f6_banka_kolon_yerlesimi(banka, ornek_yolu, tmp_path):
    e = ekstre_oku(ornek_yolu(banka))
    yol = bordro_f1f6.yaz(e.islemler, tmp_path / "f.xls", banka)
    kitap = xlrd.open_workbook(str(yol), formatting_info=True)
    ws = kitap.sheet_by_index(0)
    assert ws.name == "Sayfa"
    assert ws.row_values(0) == ["F1", "F2", "F3", "F4", "F5", "F6"]
    assert ws.nrows == len(e.islemler) + 1
    g, c, negatif = BEKLENEN_YERLESIM[banka]
    yerlesim = bordro_f1f6.profil_yukle().yerlesim(banka)
    for r, islem in enumerate(e.islemler, start=1):
        satir = ws.row_values(r)
        assert satir[0] == islem.tarih and satir[1] == ""
        if islem.tutar > 0:
            assert satir[g] == pytest.approx(islem.tutar) and satir[c] == ""
        else:
            assert satir[g] == ""
            assert satir[c] == pytest.approx(islem.tutar if negatif else -islem.tutar)
        metinler = {yerlesim.kolonlar[k]: satir[k] for k in range(6)}
        assert metinler["aciklama"] == islem.aciklama
        assert metinler["referans"] == islem.referans
    # Tutar kolonları 0.00, diğerleri @ biçiminde
    for k in range(6):
        xf = kitap.xf_list[ws.cell_xf_index(1, k)]
        bicim = kitap.format_map[xf.format_key].format_str
        assert bicim == ("0.00" if k in (g, c) else "@")


def test_f1f6_kalin_baslik(ornek_yolu, tmp_path):
    e = ekstre_oku(ornek_yolu("garanti"))
    yol = bordro_f1f6.yaz(e.islemler, tmp_path / "g.xls", "garanti")
    kitap = xlrd.open_workbook(str(yol), formatting_info=True)
    ws = kitap.sheet_by_index(0)
    kalin = {k for k in range(6) if kitap.font_list[kitap.xf_list[ws.cell_xf_index(0, k)].font_index].bold}
    assert kalin == {0, 2, 3, 4, 5}


def test_f1f6_ozel_profil(ornek_yolu, tmp_path):
    profil_yolu = tmp_path / "benim.json"
    profil_yolu.write_text(json.dumps({
        "sayfa_adi": "Banka",
        "baslik": ["A", "B", "C", "D"],
        "bankalar": {},
        "varsayilan": {"kolonlar": ["tarih", "giren", "cikan", "aciklama"], "cikan_isareti": "negatif"},
    }), encoding="utf-8")
    profil = bordro_f1f6.profil_yukle(profil_yolu)
    e = ekstre_oku(ornek_yolu("ziraat"))
    yol = bordro_f1f6.yaz(e.islemler, tmp_path / "o.xls", "ziraat", profil)
    ws = xlrd.open_workbook(str(yol)).sheet_by_index(0)
    assert ws.name == "Banka" and ws.row_values(0) == ["A", "B", "C", "D"]
    cikan = next(i for i in e.islemler if i.tutar < 0)
    satir = ws.row_values(e.islemler.index(cikan) + 1)
    assert satir[2] == pytest.approx(cikan.tutar) and satir[3] == cikan.aciklama


def test_f1f6_gecersiz_profil(tmp_path):
    yol = tmp_path / "bozuk.json"
    yol.write_text(json.dumps({"bankalar": {"ziraat": {"kolonlar": ["tarih", "x"]}}}), encoding="utf-8")
    with pytest.raises(bordro_f1f6.ProfilHatasi):
        bordro_f1f6.profil_yukle(yol)


def test_ilk_satirlar_karisik():
    islemler = [_islem(100, n) for n in range(7)] + [_islem(-50, 7), _islem(20, 8)]
    sonuc = bordro_f1f6.ilk_satirlari_karistir(islemler)
    assert {i.tutar > 0 for i in sonuc[:5]} == {True, False}
    assert sonuc[4].tutar == -50 and sonuc[7].tutar == 100
    assert len(sonuc) == len(islemler)
    # Zaten karışıksa dokunmaz
    karisik = [_islem(-1, 0)] + islemler
    assert bordro_f1f6.ilk_satirlari_karistir(karisik) == karisik
    # Ters yönlü işlem hiç yoksa dokunmaz
    tek = [_islem(5, n) for n in range(8)]
    assert bordro_f1f6.ilk_satirlari_karistir(tek) == tek


# ─────────────────────────────── CLI ───────────────────────────────

@pytest.mark.parametrize("fmt, uzanti", [("csv", ".csv"), ("xlsx", ".xlsx"), ("f1f6", ".xls")])
def test_cli_cevir_dosya_basina(fmt, uzanti, ornek_klasoru, tmp_path, capsys):
    cikti = tmp_path / "cikti"
    assert main(["cevir", str(ornek_klasoru), "--cikti", str(cikti), "--format", fmt]) == 0
    beklenen = {f"{ad.rsplit('.', 1)[0]}-{fmt}{uzanti}" for ad in uret.DOSYA_ADLARI.values()}
    assert {p.name for p in cikti.iterdir()} == beklenen
    out = capsys.readouterr().out
    assert "TOPLAM" in out and str(sum(uret.ADETLER.values())) in out


def test_cli_birlestir_csv(ornek_klasoru, tmp_path):
    cikti = tmp_path / "b"
    assert main(["cevir", str(ornek_klasoru), "--cikti", str(cikti), "--format", "csv", "--birlestir"]) == 0
    dosyalar = list(cikti.iterdir())
    assert [p.name for p in dosyalar] == ["birlesik-csv.csv"]
    with dosyalar[0].open(encoding="utf-8-sig", newline="") as f:
        satirlar = list(csv.reader(f, delimiter=";"))
    assert len(satirlar) - 1 == sum(uret.ADETLER.values())
    assert {s[1] for s in satirlar[1:]} == set(BANKALAR)


def test_cli_birlestir_xlsx(ornek_klasoru, tmp_path):
    cikti = tmp_path / "bx"
    assert main(["cevir", str(ornek_klasoru), "--cikti", str(cikti), "--format", "xlsx", "--birlestir"]) == 0
    ws = load_workbook(cikti / "birlesik-xlsx.xlsx").active
    assert ws.max_row - 1 == sum(uret.ADETLER.values())


def test_cli_birlestir_f1f6_banka_basina(ornek_klasoru, tmp_path):
    cikti = tmp_path / "bf"
    assert main(["cevir", str(ornek_klasoru), "--cikti", str(cikti), "--format", "f1f6", "--birlestir",
                 "--ilk-satirlar-karisik"]) == 0
    assert {p.name for p in cikti.iterdir()} == {f"birlesik-{b}-f1f6.xls" for b in BANKALAR}


def test_cli_tani_ve_ozet(ornek_yolu, capsys):
    assert main(["tani", str(ornek_yolu("yapikredi"))]) == 0
    out = capsys.readouterr().out
    assert "banka=yapikredi" in out and "satır=12" in out
    assert main(["ozet", str(ornek_yolu("turkiyefinans"))]) == 0
    out = capsys.readouterr().out
    assert "8 işlem" in out and "–" in out


def test_cli_olmayan_girdi(tmp_path, capsys):
    assert main(["tani", str(tmp_path / "yok.xls")]) == 2
